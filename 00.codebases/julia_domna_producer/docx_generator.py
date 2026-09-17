"""
HISTORYSNOOZE: DOCX GENERATOR
Module: docx_generator.py
SSOT-compliant with 02_SCRIPT_SOP_&_GATEKEEPERS.md.
Rule <= 150 lines strictly enforced.
"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from outline_data import JULIA_DOMNA_OUTLINE
from script_aggregator import ALL_PARTS


def generate_outline_docx(output_path: Path) -> None:
    """Generate Outline - Julia Domna.docx with professional formatting."""
    doc = Document()

    title_p = doc.add_paragraph()
    title_run = title_p.add_run(JULIA_DOMNA_OUTLINE["title"])
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run("Chronological 3-Act / 15-Part Chronicle Outline")
    sub_run.italic = True
    sub_run.font.size = Pt(12)
    doc.add_paragraph()

    for act in JULIA_DOMNA_OUTLINE["acts"]:
        act_heading = doc.add_heading(level=1)
        act_run = act_heading.add_run(f"ACT {act['act_number']}: {act['act_title'].upper()}")
        act_run.font.size = Pt(14)
        act_run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)

        for part in act["parts"]:
            p_part = doc.add_paragraph()
            p_part.paragraph_format.space_before = Pt(6)
            p_part.paragraph_format.space_after = Pt(2)
            title_run = p_part.add_run(f"Part {part['part']:02d}: {part['title']}")
            title_run.bold = True
            title_run.font.size = Pt(11)

            p_sum = doc.add_paragraph()
            p_sum.paragraph_format.left_indent = Inches(0.25)
            p_sum.paragraph_format.space_after = Pt(6)
            sum_run = p_sum.add_run(part["summary"])
            sum_run.font.size = Pt(10)

    doc.save(str(output_path))


def generate_script_docx(output_path: Path) -> None:
    """Generate Script - Julia Domna.docx with 15 parts / 150 beats."""
    doc = Document()

    title_p = doc.add_paragraph()
    title_run = title_p.add_run(JULIA_DOMNA_OUTLINE["title"])
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    sub_p = doc.add_paragraph()
    sub_run = sub_p.add_run("Canonical 15-Part Voiceover Script (150 Narrative Beats / GK2 Sanitized)")
    sub_run.italic = True
    sub_run.font.size = Pt(12)
    doc.add_paragraph()

    for part_num in range(1, 16):
        part_heading = doc.add_heading(level=1)
        p_run = part_heading.add_run(f"PART {part_num:02d}")
        p_run.font.size = Pt(14)
        p_run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)

        paragraphs = ALL_PARTS[part_num]
        for idx, para in enumerate(paragraphs, start=1):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(para)
            run.font.size = Pt(11)

    doc.save(str(output_path))
